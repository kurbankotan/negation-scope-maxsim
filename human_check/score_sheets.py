"""Scores the filled annotation workbooks of the human check against the key (revision step 2, October 2026).

usage: python score_sheets.py KEY_CSV OUT_DIR WORKBOOK1 [WORKBOOK2 ...]

For each workbook (annotator): agreement with the gold label per set (perm, fresh) and arm (O, P, N, S), the rate of
"?" answers and of "not natural" marks, scenario-level agreement (all four arms); between annotators: raw agreement
and Cohen's kappa on {A, B, ?} and on {A, B} (items both answered A or B). Lists the items on which every annotator
disagrees with the gold label (exclusion candidates for the robustness analysis) and the items marked not natural.
Writes human_check_results.json and human_check_results.md to OUT_DIR. Blank answers count as "?".
"""
import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ARMS = ("O", "P", "N", "S")


def read_key(path):
    return {(r["code"], int(r["pos"])): r for r in csv.DictReader(open(path, encoding="utf-8"))}


def read_workbook(path):
    ws = load_workbook(path, data_only=True)["Maddeler"]
    head = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[col["Senaryo"]] is None:
            continue
        code, pos = str(row[col["Senaryo"]]).strip(), int(row[col["Sıra"]])
        ans = row[col["Cevap"]]
        ans = str(ans).strip().upper() if ans not in (None, "") else "?"
        if ans not in ("A", "B", "?"):
            ans = "?"
        nat = row[col["Doğal değil"]]
        note = row[col["Not"]]
        out[(code, pos)] = {"answer": ans, "not_natural": bool(nat not in (None, "")), "note": (str(note).strip() if note not in (None, "") else "")}
    return out


def pct(a, n):
    return (100.0 * a / n) if n else None


def kappa(pairs):
    """Cohen's kappa for a list of (label1, label2)."""
    n = len(pairs)
    if n == 0:
        return None
    cats = sorted({x for p in pairs for x in p})
    po = sum(a == b for a, b in pairs) / n
    c1, c2 = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum((c1[c] / n) * (c2[c] / n) for c in cats)
    return (po - pe) / (1 - pe) if pe < 1 else 1.0


def score(key, books):
    names = list(books)
    res = {"annotators": names, "items": len(key), "per_annotator": {}, "between": {}, "exclusion_candidates": [], "not_natural": []}
    for name, b in books.items():
        missing = [k for k in key if k not in b]
        blk = {"missing_rows": len(missing), "by_set_arm": {}, "scenario_all_four": {}}
        for split in ("perm", "fresh"):
            by_scen = defaultdict(list)
            for arm in ARMS:
                ks = [k for k, r in key.items() if r["set"] == split and r["arm"] == arm]
                ans = [b.get(k, {"answer": "?", "not_natural": False})["answer"] for k in ks]
                gold = [key[k]["label"] for k in ks]
                nat = [b.get(k, {"not_natural": False})["not_natural"] for k in ks]
                agree = sum(a == g for a, g in zip(ans, gold))
                decided = [(a, g) for a, g in zip(ans, gold) if a != "?"]
                blk["by_set_arm"][f"{split}.{arm}"] = {"n": len(ks), "agree_with_gold_pct": pct(agree, len(ks)), "question_mark_pct": pct(sum(a == "?" for a in ans), len(ks)),
                                                        "agree_among_decided_pct": pct(sum(a == g for a, g in decided), len(decided)), "not_natural_pct": pct(sum(nat), len(ks))}
                for k, a, g in zip(ks, ans, gold):
                    by_scen[key[k]["scenario_id"]].append(a == g)
            blk["scenario_all_four"][split] = {"n": len(by_scen), "pct": pct(sum(all(v) for v in by_scen.values()), len(by_scen))}
            allk = [k for k, r in key.items() if r["set"] == split]
            blk["by_set_arm"][f"{split}.all"] = {"n": len(allk), "agree_with_gold_pct": pct(sum(b.get(k, {"answer": "?"})["answer"] == key[k]["label"] for k in allk), len(allk)),
                                                 "question_mark_pct": pct(sum(b.get(k, {"answer": "?"})["answer"] == "?" for k in allk), len(allk)),
                                                 "not_natural_pct": pct(sum(b.get(k, {"not_natural": False})["not_natural"] for k in allk), len(allk))}
        res["per_annotator"][name] = blk
    if len(names) >= 2:
        for i in range(len(names)):
            for j in range(i + 1, len(names)):
                a, b = books[names[i]], books[names[j]]
                for split in ("perm", "fresh", "all"):
                    ks = [k for k, r in key.items() if split == "all" or r["set"] == split]
                    pairs3 = [(a.get(k, {"answer": "?"})["answer"], b.get(k, {"answer": "?"})["answer"]) for k in ks]
                    pairs2 = [p for p in pairs3 if "?" not in p]
                    res["between"][f"{names[i]}|{names[j]}|{split}"] = {"n": len(pairs3), "raw_agreement_pct": pct(sum(x == y for x, y in pairs3), len(pairs3)),
                                                                        "kappa_ABq": kappa(pairs3), "n_both_decided": len(pairs2), "kappa_AB": kappa(pairs2)}
    for k, r in sorted(key.items()):
        answers = {n: books[n].get(k, {"answer": "?"})["answer"] for n in names}
        if names and all(v != r["label"] for v in answers.values()):
            res["exclusion_candidates"].append({"item_id": r["item_id"], "scenario_id": r["scenario_id"], "set": r["set"], "arm": r["arm"], "gold": r["label"], "answers": answers})
        nat = [n for n in names if books[n].get(k, {"not_natural": False})["not_natural"]]
        if nat:
            res["not_natural"].append({"item_id": r["item_id"], "set": r["set"], "arm": r["arm"], "by": nat, "notes": {n: books[n].get(k, {}).get("note", "") for n in nat}})
    res["summary"] = {"exclusion_candidates": len(res["exclusion_candidates"]),
                      "exclusion_candidates_by_set_arm": dict(Counter(f"{e['set']}.{e['arm']}" for e in res["exclusion_candidates"])),
                      "not_natural_items": len(res["not_natural"])}
    return res


def markdown(res):
    lines = ["# Human check of NegPairs-Perm and the NegPairs-Fresh sample", ""]
    for name, blk in res["per_annotator"].items():
        lines += [f"## {name} (missing rows: {blk['missing_rows']})", "", "| set.arm | n | agree with gold % | ? % | agree among decided % | not natural % |", "|---|---|---|---|---|---|"]
        for k, v in blk["by_set_arm"].items():
            f = lambda x: "" if x is None else f"{x:.1f}"
            lines.append(f"| {k} | {v['n']} | {f(v['agree_with_gold_pct'])} | {f(v['question_mark_pct'])} | {f(v.get('agree_among_decided_pct'))} | {f(v['not_natural_pct'])} |")
        lines.append("")
        for split, v in blk["scenario_all_four"].items():
            lines.append(f"* {split}: all four arms agree with gold in {v['pct']:.1f}% of {v['n']} scenarios")
        lines.append("")
    if res["between"]:
        lines += ["## Between annotators", "", "| pair | set | n | raw agreement % | kappa (A/B/?) | n both decided | kappa (A/B) |", "|---|---|---|---|---|---|---|"]
        for k, v in res["between"].items():
            a, b, split = k.split("|")
            lines.append(f"| {a} vs {b} | {split} | {v['n']} | {v['raw_agreement_pct']:.1f} | {v['kappa_ABq']:.3f} | {v['n_both_decided']} | "
                         f"{'' if v['kappa_AB'] is None else f'{v['kappa_AB']:.3f}'} |")
        lines.append("")
    lines += [f"Exclusion candidates (every annotator disagrees with gold): {res['summary']['exclusion_candidates']} "
              f"({res['summary']['exclusion_candidates_by_set_arm']}); items marked not natural by at least one annotator: {res['summary']['not_natural_items']}."]
    return "\n".join(lines)


if __name__ == "__main__":
    key = read_key(sys.argv[1])
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)
    books = {Path(p).stem: read_workbook(p) for p in sys.argv[3:]}
    res = score(key, books)
    (out_dir / "human_check_results.json").write_text(json.dumps(res, indent=1, ensure_ascii=False), encoding="utf-8")
    (out_dir / "human_check_results.md").write_text(markdown(res), encoding="utf-8")
    print(markdown(res))
