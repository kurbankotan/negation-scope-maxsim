"""Annotation workbooks for the human check of NegPairs-Perm (all 274 scenarios) and a random sample of NegPairs-Fresh
(60 scenarios, 15 per family), revision step 2 (October 2026).

Design. Every scenario appears as four consecutive rows, one per context (the four arms O, P, N, S in a random
order, unlabeled), with the question and the two options repeated; the annotator reads the question once and answers
A, B or ? for each context. Arm identity and the gold label are not in the workbook; they are kept in
`human_check_key.csv` next to this script. Scenarios of the two sets are mixed in one random order (seed 20261022).

usage: python make_sheets.py DATA_DIR OUT_DIR
    DATA_DIR  folder with perm.jsonl and fresh.jsonl (np_data_v3/negpairs_v3)
    OUT_DIR   where the two identical workbooks (one per annotator) and the key are written
"""
import csv
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20261022
FRESH_PER_FAMILY = 15
ARMS = ("O", "P", "N", "S")


def load(path):
    return [json.loads(l) for l in open(path, encoding="utf-8") if l.strip()]


def build(data_dir: Path, out_dir: Path):
    rng = random.Random(SEED)
    perm, fresh = load(data_dir / "perm.jsonl"), load(data_dir / "fresh.jsonl")
    by = {}
    for r in perm + fresh:
        by.setdefault(r["scenario_id"], {})[r["arm"]] = r
    perm_ids = sorted({r["scenario_id"] for r in perm})
    fresh_ids = []
    for fam in sorted({r["family"] for r in fresh}):
        ids = sorted({r["scenario_id"] for r in fresh if r["family"] == fam})
        fresh_ids.extend(rng.sample(ids, FRESH_PER_FAMILY))
    scen = [(sid, "perm") for sid in perm_ids] + [(sid, "fresh") for sid in fresh_ids]
    rng.shuffle(scen)
    rows, key = [], []
    for k, (sid, split) in enumerate(scen, start=1):
        code = f"S{k:03d}"
        arms = list(ARMS)
        rng.shuffle(arms)
        for j, arm in enumerate(arms, start=1):
            r = by[sid][arm]
            rows.append({"code": code, "pos": j, "context": r["context"], "question": r["question"], "A": r["options"][0], "B": r["options"][1]})
            key.append({"code": code, "pos": j, "set": split, "scenario_id": sid, "item_id": r["item_id"], "arm": arm, "label": "AB"[int(r["label"])],
                        "family": r["family"], "seen_family": r["seen_family"]})
    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "human_check_key.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(key[0]))
        w.writeheader()
        w.writerows(key)
    for name in ("degerlendirici1", "degerlendirici2"):
        write_workbook(rows, out_dir / f"NegPairs_insan_denetimi_{name}.xlsx", len(perm_ids), len(fresh_ids))
    return len(rows), len(perm_ids), len(fresh_ids)


def write_workbook(rows, path: Path, n_perm: int, n_fresh: int):
    wb = Workbook()
    ws = wb.active
    ws.title = "Talimat"
    font = Font(name="Arial", size=11)
    bold = Font(name="Arial", size=11, bold=True)
    title = Font(name="Arial", size=14, bold=True)
    fill = PatternFill("solid", fgColor="FFFF00")
    lines = [
        ("NegPairs insan denetimi", title),
        ("", font),
        ("Ne yapılacak", bold),
        ("Her satırda kısa bir bağlam, bir soru ve iki seçenek var. Yalnızca bağlamda yazılana göre sorunun cevabı hangi seçenekse "
         "onu işaretleyin: Cevap sütununa A veya B yazın (açılır liste). Bağlam cevabı belirlemiyorsa veya iki seçenek de "
         "savunulabilirse ? yazın.", font),
        ("Aynı senaryo kodu (S001 gibi) dört ardışık satırda tekrar eder: soru ve seçenekler aynıdır, bağlam cümlesi değişir. "
         "Her bağlamı kendi başına, önceki satıra bakmadan değerlendirin; dört satırın cevapları aynı veya farklı olabilir.", font),
        ("Bağlam cümlesi size bozuk, anlamsız veya bir İngilizce konuşurun söylemeyeceği gibi geliyorsa \"Doğal değil\" sütununa X "
         "koyun (cevabı yine verin). Not sütunu isteğe bağlıdır.", font),
        ("", font),
        ("Kurallar", bold),
        ("Dış bilgi veya tahmin kullanmayın; yalnızca bağlamın söylediği önemlidir. Sözlük veya arama kullanmayın. Diğer "
         "değerlendiriciyle görüşmeden doldurun. Satırları atlamayın; emin olmadığınızda ? seçin.", font),
        ("Soru sırası rastgeledir; hız için önce soruyu ve seçenekleri okuyup sonra dört bağlamı sırayla cevaplamak en kolayıdır.", font),
        (f"Toplam {n_perm + n_fresh} senaryo, {4 * (n_perm + n_fresh)} satır ({n_perm} senaryo ana küme, {n_fresh} senaryo ek örneklem; "
         "kodlar hangisinin hangisi olduğunu göstermez). Beklenen süre yaklaşık 3 saat; ara verebilirsiniz, dosyayı kaydedin.", font),
        ("", font),
        ("Doldurulacak hücreler", bold),
        ("Yalnızca sarı başlıklı üç sütun doldurulur: Cevap (zorunlu: A, B veya ?), Doğal değil (isteğe bağlı: X), Not (isteğe bağlı). "
         "Diğer hücreleri değiştirmeyin, satır eklemeyin veya silmeyin.", font),
        ("", font),
        ("Örnek", bold),
        ("Bağlam: \"The café is open on Sundays, but the bakery is closed on Sundays.\"  "
         "Soru: \"Where can she buy bread on Sunday?\"  A: At the bakery  B: At the café  →  Cevap: B", font),
        ("Bağlam: \"The café is closed on Sundays, but the bakery is open on Sundays.\"  (aynı soru ve seçenekler)  →  Cevap: A", font),
        ("Bağlam: \"The café and the bakery are both open on Sundays.\"  (aynı soru ve seçenekler)  →  Cevap: ?", font),
        ("", font),
        ("English summary", bold),
        ("For each row, read the context, the question and the two options and mark in the Answer column which option the "
         "context supports (A or B; ? if the context does not decide). The same scenario code repeats on four consecutive rows "
         "with different contexts: judge each context on its own. Put X in the \"Not natural\" column if the sentence is "
         "ungrammatical or unidiomatic. Use only the context, no outside knowledge; do not confer with the other annotator; "
         "fill the three yellow columns only.", font),
    ]
    for i, (text, f) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    ws2 = wb.create_sheet("Maddeler")
    head = ["Senaryo", "Sıra", "Bağlam", "Soru", "Seçenek A", "Seçenek B", "Cevap", "Doğal değil", "Not"]
    widths = [10, 6, 70, 42, 26, 26, 9, 11, 30]
    thin = Side(style="thin", color="999999")
    for j, (h, w) in enumerate(zip(head, widths), start=1):
        c = ws2.cell(row=1, column=j, value=h)
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = Border(bottom=thin)
        if h in ("Cevap", "Doğal değil", "Not"):
            c.fill = fill
        ws2.column_dimensions[get_column_letter(j)].width = w
    for i, r in enumerate(rows, start=2):
        vals = [r["code"], r["pos"], r["context"], r["question"], r["A"], r["B"], None, None, None]
        for j, v in enumerate(vals, start=1):
            c = ws2.cell(row=i, column=j, value=v)
            c.font = font
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if r["pos"] == 1 and i > 2:
            for j in range(1, len(head) + 1):
                ws2.cell(row=i, column=j).border = Border(top=thin)
    n = len(rows) + 1
    dv = DataValidation(type="list", formula1='"A,B,?"', allow_blank=True, showErrorMessage=True, errorTitle="Cevap", error="A, B veya ? yazın")
    dv.add(f"G2:G{n}")
    dv2 = DataValidation(type="list", formula1='"X"', allow_blank=True, showErrorMessage=True, errorTitle="Doğal değil", error="X yazın veya boş bırakın")
    dv2.add(f"H2:H{n}")
    ws2.add_data_validation(dv)
    ws2.add_data_validation(dv2)
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(head))}{n}"
    wb.save(path)


if __name__ == "__main__":
    data_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[2] / "np_data_v3" / "negpairs_v3"
    out_dir = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).parent / "sheets"
    n_rows, n_perm, n_fresh = build(data_dir, out_dir)
    print(f"{n_rows} rows ({n_perm} Perm + {n_fresh} Fresh scenarios) -> {out_dir}")
